"""Импорт расписания колледжа (XLS/XLSX) для группы ИС-21.

Что делает парсер:
  1) читает все листы, разворачивает объединённые ячейки;
  2) ищет столбец нужной группы по заголовкам (ИС-21, ИС 21, ис-21 …);
  3) ограничивает чтение ТОЛЬКО столбцами этой группы — соседние группы
     (ИС-22, ЭК-31) в данные не попадают;
  4) определяет дату: параметр импорта → дата в файле → имя листа → имя файла;
  5) пустые ячейки пропускает, а не выдумывает занятия;
  6) если дату/занятия определить не удалось — возвращает предупреждения
     и пустой список (никаких «придуманных» пар).

Поддержаны форматы:
  * «блочный»: дата (день) + строки занятий под ней в одном листе;
  * «одна дата на лист»: дата в шапке листа или в имени листа/файла;
  * несколько листов в одной книге.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

GROUP_DEFAULT = "ИС-21"

# --------------------------------------------------------------------------- #
# Абстракция книги
# --------------------------------------------------------------------------- #
@dataclass
class MergedRange:
    min_row: int
    max_row: int
    min_col: int
    max_col: int


@dataclass
class SheetGrid:
    name: str
    rows: list[list[str]]
    merged: list[MergedRange] = field(default_factory=list)

    def cell(self, row: int, col: int) -> str:
        if row < 1 or col < 1 or row > len(self.rows):
            return ""
        line = self.rows[row - 1]
        if col > len(line):
            return ""
        return (line[col - 1] or "").strip()

    @property
    def n_rows(self) -> int:
        return len(self.rows)

    @property
    def n_cols(self) -> int:
        return max((len(r) for r in self.rows), default=0)

    def merge_extent(self, row: int, col: int) -> tuple[int, int]:
        for rng in self.merged:
            if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
                return rng.min_col, rng.max_col
        return col, col


def _s(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\xa0", " ").strip()


def load_workbook_bytes(data: bytes, filename: str = "") -> list[SheetGrid]:
    """Прочитать XLSX (openpyxl) или XLS (xlrd) в нормализованную сетку."""
    name = (filename or "").lower()
    if name.endswith(".xls") and not name.endswith(".xlsx"):
        return _load_xls(data)
    try:
        return _load_xlsx(data)
    except Exception as exc:
        raise ValueError(f"Не удалось прочитать файл как XLSX: {exc}") from exc


def _load_xlsx(data: bytes) -> list[SheetGrid]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True)
    sheets: list[SheetGrid] = []
    for ws in wb.worksheets:
        merged: list[MergedRange] = []
        grid = [[_s(c) for c in row] for row in ws.iter_rows(values_only=True)]
        for rng in ws.merged_cells.ranges:
            merged.append(MergedRange(rng.min_row, rng.max_row, rng.min_col, rng.max_col))
            value = _s(ws.cell(row=rng.min_row, column=rng.min_col).value)
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    while len(grid) < r:
                        grid.append([])
                    while len(grid[r - 1]) < c:
                        grid[r - 1].append("")
                    if not grid[r - 1][c - 1]:
                        grid[r - 1][c - 1] = value
        sheets.append(SheetGrid(name=ws.title, rows=grid, merged=merged))
    return sheets


def _load_xls(data: bytes) -> list[SheetGrid]:
    import xlrd

    book = xlrd.open_workbook(file_contents=data)
    sheets: list[SheetGrid] = []
    for sh in book.sheets():
        rows = [[_s(sh.cell_value(r, c)) for c in range(sh.ncols)] for r in range(sh.nrows)]
        merged = []
        for rlo, rhi, clo, chi in getattr(sh, "merged_cells", []):
            merged.append(MergedRange(rlo + 1, rhi, clo + 1, chi))
            value = rows[rlo][clo] if rlo < len(rows) and clo < len(rows[rlo]) else ""
            for r in range(rlo, min(rhi, len(rows))):
                for c in range(clo, min(chi, len(rows[r]))):
                    if not rows[r][c]:
                        rows[r][c] = value
        sheets.append(SheetGrid(name=sh.name, rows=rows, merged=merged))
    return sheets


# --------------------------------------------------------------------------- #
# Поиск столбца группы
# --------------------------------------------------------------------------- #
GROUP_RE = re.compile(r"([А-ЯЁA-Z]{2,5})\s*[-–—]?\s*(\d{1,2})", re.IGNORECASE)


def normalize_group(text: str) -> str:
    return re.sub(r"[\s\-–—]", "", (text or "").upper().replace("Ё", "Е"))


def groups_in_cell(text: str) -> list[tuple[str, int]]:
    """Все пары (буквы, номер) группы в ячейке."""
    found = []
    for letters, number in GROUP_RE.findall((text or "").upper()):
        found.append((letters.replace("Ё", "Е"), int(number)))
    return found


def find_group_column(sheet: SheetGrid, group: str) -> tuple[int, int, int] | None:
    """Найти (row, col, max_col) заголовка группы. max_col — граница merge."""
    target = normalize_group(group)
    if not target:
        return None
    for r in range(1, min(sheet.n_rows, 40) + 1):
        for c in range(1, sheet.n_cols + 1):
            text = sheet.cell(r, c)
            if not text or len(text) > 40:
                continue
            for letters, number in groups_in_cell(text):
                if normalize_group(f"{letters}{number}") == target:
                    lo, hi = sheet.merge_extent(r, c)
                    return r, c, hi
    return None


# --------------------------------------------------------------------------- #
# Время, дата, разбор занятия
# --------------------------------------------------------------------------- #
TIME_RANGE_RE = re.compile(r"(\d{1,2})[:.\s]?(\d{2})\s*[-–—]\s*(\d{1,2})[:.\s]?(\d{2})")
DATE_RE = re.compile(r"\b(\d{1,2})[.\/](\d{1,2})(?:[.\/](\d{2,4}))?\b")
RU_MONTHS = {
    "январ": 1, "феврал": 2, "март": 3, "апрел": 4, "ма": 5, "июн": 6,
    "июл": 7, "август": 8, "сентябр": 9, "октябр": 10, "ноябр": 11, "декабр": 12,
}
LESSON_TYPES = ("лекц", "практ", "лабор", "семин", "контрольн", "зачёт", "зачет",
                "экзамен", "консультац", "самостоятельн")
WEEKDAYS = ("понедельник", "вторник", "среда", "четверг", "пятница", "суббота", "воскресенье")
ROOM_RE = re.compile(
    r"(?:каб(?:инет)?|ауд(?:итория)?|лаб(?:оратория)?|№)\b\s*\.?\s*"
    r"([0-9]{1,4}[А-Яа-яA-Za-z]?(?:[/-][0-9А-Яа-я]{1,4})?)",
    re.IGNORECASE,
)
INITIALS_RE = re.compile(r"[А-ЯЁ]\.\s?(?:[А-ЯЁ]\.)?")
CAPWORD_RE = re.compile(r"[А-ЯЁ][а-яё\-]+")
FULLNAME_RE = re.compile(r"[А-ЯЁ][а-яё\-]+\s+[А-ЯЁ][а-яё\-]+\s+[А-ЯЁ][а-яё\-]+")


def parse_date(text: str, default_year: int | None = None) -> date | None:
    """Распознать дату в строке (12.05.2025, 12.05, «12 мая 2025»)."""
    if not text:
        return None
    raw = str(text).strip()
    m = DATE_RE.search(raw)
    if m:
        day, month = int(m.group(1)), int(m.group(2))
        year = int(m.group(3)) if m.group(3) else (default_year or datetime.now().year)
        if year < 100:
            year += 2000
        try:
            return date(year, month, day)
        except ValueError:
            pass
    low = raw.lower()
    dm = re.search(r"(\d{1,2})\s+([а-яё]+)", low)
    if dm:
        day = int(dm.group(1))
        for stem, month in RU_MONTHS.items():
            if dm.group(2).startswith(stem):
                year = default_year or datetime.now().year
                ym = re.search(r"(20\d{2})", low)
                if ym:
                    year = int(ym.group(1))
                try:
                    return date(year, month, day)
                except ValueError:
                    return None
    return None


def parse_time_range(text: str) -> tuple[str, str] | None:
    m = TIME_RANGE_RE.search(text or "")
    if not m:
        return None
    h1, m1, h2, m2 = (int(g) for g in m.groups())
    if not (0 <= h1 <= 23 and 0 <= h2 <= 24 and 0 <= m1 < 60 and 0 <= m2 < 60):
        return None
    return f"{h1:02d}:{m1:02d}", f"{h2:02d}:{m2:02d}"


def extract_teacher(raw: str) -> tuple[str, str]:
    """Вытащить ФИО преподавателя и вернуть (ФИО, остаток текста)."""
    work = raw
    for _ in range(2):
        m = INITIALS_RE.search(work)
        if not m:
            break
        before = work[: m.start()].rstrip(" ,;")
        after = work[m.end():].lstrip(" ,;")
        before_matches = list(CAPWORD_RE.finditer(before))
        before_word = before_matches[-1] if before_matches else None
        after_word = CAPWORD_RE.match(after)
        teacher = ""
        if before_word and before_word.start() > 0:
            teacher = f"{before_word.group(0)} {m.group(0).strip()}"
            work = (before[: before_word.start()].rstrip(" ,;") + " " + after).strip(" ,;")
        elif after_word:
            teacher = f"{m.group(0).strip()} {after_word.group(0)}"
            work = (before + " " + after[after_word.end():]).strip(" ,;")
        elif before_word:
            teacher = f"{before_word.group(0)} {m.group(0).strip()}"
            work = (before[: before_word.start()].rstrip(" ,;") + " " + after).strip(" ,;")
        if teacher:
            return teacher, work
    full = FULLNAME_RE.search(work)
    if full:
        teacher = full.group(0)
        return teacher, (work[: full.start()] + " " + work[full.end():]).strip(" ,;")
    return "", work


def split_lesson(text: str) -> dict:
    """Разбить содержимое ячейки на предмет / преподавателя / кабинет / тип."""
    raw = " ".join((text or "").split())
    room = ""
    rm = ROOM_RE.search(raw)
    if rm:
        room = rm.group(0).strip(" ,;")
        raw = (raw[: rm.start()] + " " + raw[rm.end():]).strip(" ,;")

    teacher, raw = extract_teacher(raw)

    lesson_type = ""
    low = raw.lower()
    for keyword in LESSON_TYPES:
        if keyword in low:
            lesson_type = keyword
            break

    # скобки с типом занятия: «Физика (лабораторная)» → предмет «Физика»
    paren = re.search(r"\(([^)]*)\)", raw)
    if paren:
        inner = paren.group(1).lower()
        if any(k in inner for k in LESSON_TYPES) or len(inner) <= 24:
            raw = (raw[: paren.start()] + " " + raw[paren.end():]).strip(" ,;")

    raw = re.sub(r"\b(лекц\w*|практ\w*|лабор\w*|семин\w*|контрольн\w*|зач[её]т|экзамен\w*|консультац\w*|самостоятельн\w*)\b",
                 " ", raw, flags=re.IGNORECASE)
    subject = re.sub(r"\s{2,}", " ", raw).strip(" ,;.-—–")
    return {"subject": subject, "teacher": teacher, "room": room, "lesson_type": lesson_type}


# --------------------------------------------------------------------------- #
# Импорт
# --------------------------------------------------------------------------- #
@dataclass
class ImportResult:
    group: str
    entries: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    sheets_scanned: int = 0
    sheets_matched: int = 0

    def as_dict(self) -> dict:
        return {
            "group": self.group,
            "entries": self.entries,
            "warnings": self.warnings,
            "lessons": len(self.entries),
            "sheets_scanned": self.sheets_scanned,
            "sheets_matched": self.sheets_matched,
        }


DATE_WITH_YEAR_RE = re.compile(r"(\d{1,2})[._\-/](\d{1,2})[._\-/](\d{4})")


def _parse_date_with_year(text: str) -> date | None:
    """Дата только в полном виде (дд.мм.гггг) — для имён файлов и листов."""
    match = DATE_WITH_YEAR_RE.search(text or "")
    if not match:
        return None
    try:
        return date(int(match.group(3)), int(match.group(2)), int(match.group(1)))
    except ValueError:
        return None


def _sheet_default_date(sheet: SheetGrid, filename: str, override: date | None) -> tuple[date | None, str]:
    if override:
        return override, "параметр импорта"
    for r in range(1, min(sheet.n_rows, 6) + 1):
        for c in range(1, min(sheet.n_cols, 8) + 1):
            parsed = parse_date(sheet.cell(r, c))
            if parsed:
                return parsed, f"ячейка {sheet.name}!{r},{c}"
    sheet_date = parse_date(sheet.name) or _parse_date_with_year(sheet.name)
    if sheet_date:
        return sheet_date, "имя листа"
    file_date = _parse_date_with_year(Path(filename).stem)
    if file_date:
        return file_date, "имя файла"
    return None, ""


def parse_schedule(data: bytes, filename: str = "", group: str = GROUP_DEFAULT,
                   date_override: str | None = None) -> ImportResult:
    """Разобрать файл расписания. Ничего не выдумывает: пусто → warnings."""
    result = ImportResult(group=group)
    override = parse_date(date_override) if date_override else None
    if date_override and not override:
        result.warnings.append(f"Дата «{date_override}» не распознана — использованы даты из файла")

    sheets = load_workbook_bytes(data, filename)
    if not sheets:
        result.warnings.append("Файл не содержит листов")
        return result

    for sheet in sheets:
        result.sheets_scanned += 1
        location = find_group_column(sheet, group)
        if location is None:
            result.warnings.append(f"Лист «{sheet.name}»: столбец группы {group} не найден — лист пропущен")
            continue
        header_row, col_from, col_to = location
        result.sheets_matched += 1
        base_date, date_source = _sheet_default_date(sheet, filename, override)
        current_date = base_date
        block_date_source = date_source

        if base_date is None:
            # даты нет на уровне листа — возможно, «блочный» формат: ищем даты в строках
            block_date_source = "строка с датой"

        slot = 0
        for r in range(header_row + 1, sheet.n_rows + 1):
            row_text = " ".join(sheet.cell(r, c) for c in range(1, min(sheet.n_cols, col_from - 1) + 1))
            row_date = parse_date(row_text) if base_date is None or _looks_like_date_row(sheet, r) else None
            if row_date:
                current_date = row_date
                block_date_source = f"строка {r}"
                slot = 0
                continue

            if _is_header_like(sheet, r, col_from):
                continue
            if normalize_group(sheet.cell(r, col_from)) == normalize_group(group):
                continue  # повтор заголовка группы

            cell_text = ""
            for c in range(col_from, col_to + 1):
                value = sheet.cell(r, c)
                if value:
                    cell_text = value
                    break
            if not cell_text:
                continue

            # время пары: ищем в левой части строки
            times = None
            for c in range(1, max(col_from - 1, 1) + 1):
                times = parse_time_range(sheet.cell(r, c))
                if times:
                    break
            if times is None:
                times = parse_time_range(cell_text)
            slot += 1

            if current_date is None:
                result.warnings.append(
                    f"Лист «{sheet.name}», строка {r}: дата занятия не определена — занятие пропущено"
                )
                continue

            parts = split_lesson(cell_text)
            if not parts["subject"]:
                continue
            result.entries.append({
                "group_name": group,
                "date": current_date.isoformat(),
                "slot": slot,
                "start_time": times[0] if times else "",
                "end_time": times[1] if times else "",
                "subject": parts["subject"],
                "teacher": parts["teacher"],
                "room": parts["room"],
                "lesson_type": parts["lesson_type"],
                "source_file": filename,
                "source_sheet": sheet.name,
                "date_source": block_date_source,
            })

    # устранение полных дублей (merge по вертикали)
    seen = set()
    unique = []
    for entry in result.entries:
        key = (entry["date"], entry["start_time"], entry["subject"], entry["room"], entry["teacher"])
        if key in seen:
            continue
        seen.add(key)
        unique.append(entry)
    result.entries = unique

    if not result.entries:
        result.warnings.append("Занятия не найдены: файл повреждён или формат не распознан")
    return result


def _looks_like_date_row(sheet: SheetGrid, row: int) -> bool:
    """Строка-заголовок дня: дата + день недели / короткая строка."""
    left = " ".join(sheet.cell(row, c) for c in range(1, min(sheet.n_cols, 8) + 1)).lower()
    if any(day in left for day in WEEKDAYS):
        return True
    return parse_date(left) is not None and len(left) < 40


def _is_header_like(sheet: SheetGrid, row: int, group_col: int) -> bool:
    text = sheet.cell(row, group_col).lower()
    if not text:
        return False
    return text in ("группа", "группы", "день", "дни недели", "пара", "время") or text.startswith("группа ")
